import streamlit as st
import pandas as pd
from openpyxl import load_workbook
import math
import re
import io
import openai
import zipfile
import json
from lxml import etree
from fpdf import FPDF

# ──────────────────────────────────────────────
# 1. Page Setting
# ──────────────────────────────────────────────
st.set_page_config(page_title="IS 2010 Scoring System", layout="wide")
st.title("IS 2010 Excel Guided Lab Scoring System")

# ──────────────────────────────────────────────
# 2. OpenAI Client
# ──────────────────────────────────────────────
try:
    client = openai.OpenAI(api_key=st.secrets["OPENAI_API_KEY"])
except Exception:
    client = None
    st.error("⚠️ OpenAI API Key not found. Please check your .streamlit/secrets.toml")

# ──────────────────────────────────────────────
# 3. Constants
# ──────────────────────────────────────────────
STANDARD_COLORS = {
    "Dark Red":   "#C00000",
    "Red":        "#FF0000",
    "Orange":     "#FFC000",
    "Yellow":     "#FFFF00",
    "Light Green":"#92D050",
    "Green":      "#00B050",
    "Light Blue": "#00B0F0",
    "Blue":       "#0070C0",
    "Dark Blue":  "#002060",
    "Purple":     "#7030A0",
}
SCORING_MODES = ["value_only", "format_only", "both"]
SCORING_LABELS = {"value_only": "Value Only", "format_only": "Format Only", "both": "Both (Value + Format)"}

# ──────────────────────────────────────────────
# 4. Helper: Display formatting
# ──────────────────────────────────────────────
def format_ans(f, v):
    f_str = str(f).strip() if f is not None else ""
    v_str = str(v).strip() if v is not None else ""
    if f_str.startswith('='):
        return f"{f_str} ({v_str})"
    return v_str if v_str else "Empty"

# ──────────────────────────────────────────────
# 5. Core Grading: Value Comparison
# ──────────────────────────────────────────────
def check_value_equivalence(prof_f, stud_f, prof_v, stud_v):
    if prof_v is None and stud_v is None:
        pf_str = str(prof_f).replace(" ", "").upper() if prof_f else ""
        sf_str = str(stud_f).replace(" ", "").upper() if stud_f else ""
        return bool(pf_str) and pf_str == sf_str

    try:
        v_p = float(prof_v) if prof_v is not None else 0.0
        v_s = float(stud_v) if stud_v is not None else 0.0
        values_match = math.isclose(v_p, v_s, rel_tol=1e-9, abs_tol=1e-9)
    except (ValueError, TypeError):
        values_match = str(prof_v).strip().upper() == str(stud_v).strip().upper()

    if values_match:
        p_is_formula = str(prof_f).startswith('=') if prof_f else False
        s_is_formula = str(stud_f).startswith('=') if stud_f else False
        if p_is_formula and s_is_formula:
            return str(prof_f).replace(" ", "").upper() == str(stud_f).replace(" ", "").upper()
        return True
    return False

# ──────────────────────────────────────────────
# 6. Core Grading: Format Comparison
# ──────────────────────────────────────────────
def check_format_equivalence(prof_cell, stud_cell, fmt_options=None, p_wb=None, s_wb=None):
    if fmt_options is None:
        fmt_options = {k: True for k in [
            "number_format", "font_name", "font_bold", "font_color", "font_size",
            "border", "alignment", "wrap_text", "cell_size", "merge"
        ]}

    issues = []

    if fmt_options.get("number_format"):
        p_nf = prof_cell.number_format or "General"
        s_nf = stud_cell.number_format or "General"
        if p_nf.strip() != s_nf.strip():
            issues.append(f"Number format mismatch: yours='{s_nf}' expected='{p_nf}'")

    p_font = prof_cell.font
    s_font = stud_cell.font
    if p_font and s_font:
        if fmt_options.get("font_name"):
            p_fn = p_font.name or "Calibri"
            s_fn = s_font.name or "Calibri"
            if p_fn != s_fn:
                issues.append(f"Font name mismatch: yours='{s_fn}' expected='{p_fn}'")
        if fmt_options.get("font_bold"):
            if bool(p_font.bold) != bool(s_font.bold):
                issues.append(f"Bold mismatch: yours={'Bold' if s_font.bold else 'Normal'} expected={'Bold' if p_font.bold else 'Normal'}")
        if fmt_options.get("font_color"):
            p_fc = str(p_font.color.rgb) if p_font.color and p_font.color.type == 'rgb' else "000000"
            s_fc = str(s_font.color.rgb) if s_font.color and s_font.color.type == 'rgb' else "000000"
            if p_fc[-6:] != s_fc[-6:]:
                issues.append(f"Font color mismatch: yours=#{s_fc[-6:]} expected=#{p_fc[-6:]}")
        if fmt_options.get("font_size"):
            p_fs = p_font.size or 11
            s_fs = s_font.size or 11
            if abs(float(p_fs) - float(s_fs)) > 0.5:
                issues.append(f"Font size mismatch: yours={s_fs} expected={p_fs}")

    if fmt_options.get("border"):
        p_border = prof_cell.border
        s_border = stud_cell.border
        if p_border and s_border:
            for side in ['left', 'right', 'top', 'bottom']:
                p_style = getattr(p_border, side).border_style if getattr(p_border, side) else None
                s_style = getattr(s_border, side).border_style if getattr(s_border, side) else None
                if p_style != s_style:
                    issues.append(f"Border {side} mismatch: yours='{s_style}' expected='{p_style}'")

    if fmt_options.get("alignment"):
        p_al = prof_cell.alignment
        s_al = stud_cell.alignment
        if p_al and s_al:
            p_h = p_al.horizontal or "general"
            s_h = s_al.horizontal or "general"
            p_v = p_al.vertical or "bottom"
            s_v = s_al.vertical or "bottom"
            if p_h != s_h:
                issues.append(f"Horizontal alignment mismatch: yours='{s_h}' expected='{p_h}'")
            if p_v != s_v:
                issues.append(f"Vertical alignment mismatch: yours='{s_v}' expected='{p_v}'")

    if fmt_options.get("wrap_text"):
        p_al = prof_cell.alignment
        s_al = stud_cell.alignment
        p_wrap = bool(p_al.wrap_text) if p_al else False
        s_wrap = bool(s_al.wrap_text) if s_al else False
        if p_wrap != s_wrap:
            issues.append(f"Wrap text mismatch: yours={'On' if s_wrap else 'Off'} expected={'On' if p_wrap else 'Off'}")

    if fmt_options.get("cell_size") and p_wb and s_wb:
        try:
            p_ws = p_wb[prof_cell.parent.title]
            s_ws = s_wb[stud_cell.parent.title]
            row_idx    = prof_cell.row
            col_letter = prof_cell.column_letter
            p_rh = p_ws.row_dimensions[row_idx].height or 15.0
            s_rh = s_ws.row_dimensions[row_idx].height or 15.0
            if abs(float(p_rh) - float(s_rh)) > 1.0:
                issues.append(f"Row height mismatch: yours={s_rh:.1f} expected={p_rh:.1f}")
            p_cw = p_ws.column_dimensions[col_letter].width or 8.43
            s_cw = s_ws.column_dimensions[col_letter].width or 8.43
            if abs(float(p_cw) - float(s_cw)) > 1.0:
                issues.append(f"Column width mismatch: yours={s_cw:.1f} expected={p_cw:.1f}")
        except Exception:
            pass

    if fmt_options.get("merge") and p_wb and s_wb:
        try:
            p_ws = p_wb[prof_cell.parent.title]
            s_ws = s_wb[stud_cell.parent.title]
            coord = prof_cell.coordinate
            p_merged = next((str(r) for r in p_ws.merged_cells.ranges if coord in r), None)
            s_merged = next((str(r) for r in s_ws.merged_cells.ranges if coord in r), None)
            if p_merged != s_merged:
                issues.append(f"Merge mismatch: yours='{s_merged or 'not merged'}' expected='{p_merged or 'not merged'}'")
        except Exception:
            pass

    return (len(issues) == 0), issues

# ──────────────────────────────────────────────
# 7. Sparkline Check
# ──────────────────────────────────────────────
def build_sheet_xml_map(zip_bytes):
    sheet_map = {}
    try:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as z:
            wb_xml = z.read('xl/workbook.xml').decode('utf-8', errors='replace')
            wb_tree = etree.fromstring(wb_xml.encode())
            WB_NS = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            sheets_elem = wb_tree.findall('.//main:sheet', WB_NS)
            rels_xml = z.read('xl/_rels/workbook.xml.rels').decode('utf-8', errors='replace')
            rels_tree = etree.fromstring(rels_xml.encode())
            RELS_NS = {'r': 'http://schemas.openxmlformats.org/package/2006/relationships'}
            rid_to_path = {}
            for rel in rels_tree.findall('r:Relationship', RELS_NS):
                rid  = rel.get('Id')
                path = rel.get('Target')
                if not path.startswith('xl/'):
                    path = f"xl/{path}"
                rid_to_path[rid] = path
            for sh in sheets_elem:
                name = sh.get('name')
                rid  = sh.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                if name and rid and rid in rid_to_path:
                    sheet_map[name] = rid_to_path[rid]
    except Exception as e:
        st.warning(f"[Sheet map error] {e}")
    return sheet_map


def build_xml_cache(zip_bytes, sheet_map):
    cache = {}
    try:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as z:
            for sheet_name, xml_path in sheet_map.items():
                if xml_path in z.namelist():
                    cache[sheet_name] = z.read(xml_path).decode('utf-8', errors='replace')
    except Exception as e:
        st.warning(f"[XML cache error] {e}")
    return cache


def check_sparkline_advanced(p_cache, s_cache, cell_coord, sheet_name=None):
    def extract_xml_info(xml_cache, target_cell, sn):
        try:
            clean_cell = target_cell.replace('$', '').upper()
            NS = {
                "x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
                "xm":  "http://schemas.microsoft.com/office/excel/2006/main",
            }
            contents = [xml_cache[sn]] if sn and sn in xml_cache else list(xml_cache.values())
            for xml_content in contents:
                tree = etree.fromstring(xml_content.encode())
                for sp in tree.findall('.//x14:sparkline', NS):
                    sq_elem = sp.find('xm:sqref', NS)
                    f_elem  = sp.find('xm:f',     NS)
                    if sq_elem is None or f_elem is None:
                        continue
                    norm_sq = sq_elem.text.replace('$', '').upper()
                    if norm_sq == clean_cell:
                        grp = sp.getparent().getparent()
                        return {
                            "range":       f_elem.text.strip(),
                            "type":        grp.get("type", "line"),
                            "markers":     grp.get("markers",    "0") == "1",
                            "high_point":  grp.get("highPoint",  "0") == "1",
                            "low_point":   grp.get("lowPoint",   "0") == "1",
                            "first_point": grp.get("firstPoint", "0") == "1",
                            "last_point":  grp.get("lastPoint",  "0") == "1",
                            "negative":    grp.get("negative",   "0") == "1",
                        }
            return None
        except Exception as e:
            st.warning(f"[Sparkline Error] {e}")
            return None

    p_info = extract_xml_info(p_cache, cell_coord, sheet_name)
    s_info = extract_xml_info(s_cache, cell_coord, sheet_name)
    if not p_info: return "Skip", None, None, None
    if not s_info: return False, "Missing", "Sparkline Object", "Sparkline is missing."

    p_src = p_info['range'].replace("$", "").replace(" ", "").upper()
    s_src = s_info['range'].replace("$", "").replace(" ", "").upper()
    if p_src != s_src:
        return False, f"Range: {s_info['range']}", f"Range: {p_info['range']}", "Data range error"
    if p_info['type'] != s_info['type']:
        return False, f"Type: {s_info['type']}", f"Type: {p_info['type']}", "Sparkline type error"
    marker_keys = ['markers', 'high_point', 'low_point', 'first_point', 'last_point', 'negative']
    if any(p_info[k] != s_info[k] for k in marker_keys):
        return False, "Marker settings differ", "See professor's file", "Marker configuration error"
    return True, None, None, None

# ──────────────────────────────────────────────
# 8. Unified Grading Dispatcher
# ──────────────────────────────────────────────
def grade_cell(mode, p_wb_f, p_wb_v, s_wb_f, s_wb_v, p_cache, s_cache, sn, c, fmt_options=None):
    is_sl, s_ans, p_ans, msg = check_sparkline_advanced(p_cache, s_cache, c, sheet_name=sn)
    if is_sl != "Skip":
        return (bool(is_sl), f"[Sparkline] {s_ans}", f"[Sparkline] {p_ans}", [msg] if msg else [])

    pf = p_wb_f[sn][c].value
    sf = s_wb_f[sn][c].value
    pv = p_wb_v[sn][c].value
    sv = s_wb_v[sn][c].value
    p_cell_fmt = p_wb_f[sn][c]
    s_cell_fmt = s_wb_f[sn][c]

    if mode == "value_only":
        ok = check_value_equivalence(pf, sf, pv, sv)
        issues = [] if ok else [f"Value mismatch: yours='{format_ans(sf, sv)}' expected='{format_ans(pf, pv)}'"]
        return ok, format_ans(sf, sv), format_ans(pf, pv), issues
    elif mode == "format_only":
        ok, issues = check_format_equivalence(p_cell_fmt, s_cell_fmt, fmt_options=fmt_options, p_wb=p_wb_f, s_wb=s_wb_f)
        return ok, format_ans(sf, sv), format_ans(pf, pv), issues
    elif mode == "both":
        v_ok = check_value_equivalence(pf, sf, pv, sv)
        f_ok, f_issues = check_format_equivalence(p_cell_fmt, s_cell_fmt, fmt_options=fmt_options, p_wb=p_wb_f, s_wb=s_wb_f)
        issues = []
        if not v_ok:
            issues.append(f"Value mismatch: yours='{format_ans(sf, sv)}' expected='{format_ans(pf, pv)}'")
        issues.extend(f_issues)
        return (v_ok and f_ok), format_ans(sf, sv), format_ans(pf, pv), issues
    return False, "N/A", "N/A", ["Unknown mode"]

# ──────────────────────────────────────────────
# 9. AI Feedback
# ──────────────────────────────────────────────
def get_ai_feedback(prof_ans_str, stud_ans_str, issues=None, rubric=None, custom_msg=None):
    if not client: return "AI feedback disabled."
    if custom_msg: return custom_msg

    issue_text = "\n".join(issues) if issues else "General error"
    rubric_text = json.dumps(rubric, ensure_ascii=False, indent=2) if rubric else "No rubric available."

    prompt = f"""You are a professional Excel instructor giving diagnostic feedback.

Student's Answer: {stud_ans_str}
Professor's Correct Answer: {prof_ans_str}
Issues Detected: {issue_text}
Rubric Context: {rubric_text}

Write concise diagnostic feedback (under 150 words) explaining:
1. What went wrong
2. Which rubric criterion was missed (if applicable)
3. How to fix it"""

    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": "You are a professional Excel instructor."},
                {"role": "user",   "content": prompt}
            ]
        )
        return response.choices[0].message.content
    except Exception as e:
        return f"AI Error: {e}"

# ──────────────────────────────────────────────
# 10. AI Rubric Auto-Generation
# ──────────────────────────────────────────────
def extract_sparkline_info_for_cell(p_cache, cell_coord, sheet_name=None):
    try:
        clean_cell = cell_coord.replace('$', '').upper()
        NS = {
            "x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
            "xm":  "http://schemas.microsoft.com/office/excel/2006/main",
        }
        contents = [p_cache[sheet_name]] if sheet_name and sheet_name in p_cache else list(p_cache.values())
        for xml_content in contents:
            tree = etree.fromstring(xml_content.encode())
            for sp in tree.findall('.//x14:sparkline', NS):
                sq_elem = sp.find('xm:sqref', NS)
                f_elem  = sp.find('xm:f',     NS)
                if sq_elem is None or f_elem is None:
                    continue
                if sq_elem.text.replace('$', '').upper() == clean_cell:
                    grp = sp.getparent().getparent()
                    return {
                        "type":        grp.get("type", "line"),
                        "range":       f_elem.text.strip(),
                        "markers":     grp.get("markers",    "0") == "1",
                        "high_point":  grp.get("highPoint",  "0") == "1",
                        "low_point":   grp.get("lowPoint",   "0") == "1",
                        "first_point": grp.get("firstPoint", "0") == "1",
                        "last_point":  grp.get("lastPoint",  "0") == "1",
                        "negative":    grp.get("negative",   "0") == "1",
                    }
    except Exception:
        pass
    return None


def generate_rubric(check_map, p_wb_f, p_wb_v, p_cache=None):
    if not client:
        return {}, "", []

    cell_data = []
    for sn, cells in check_map.items():
        for c in cells:
            try:
                f = p_wb_f[sn][c].value
                v = p_wb_v[sn][c].value
                sl_info = extract_sparkline_info_for_cell(p_cache, c, sheet_name=sn) if p_cache else None
                if sl_info:
                    active_markers = [k for k in ["markers", "high_point", "low_point",
                                                   "first_point", "last_point", "negative"]
                                      if sl_info.get(k)]
                    cell_data.append({
                        "sheet": sn, "cell": c, "type": "sparkline",
                        "formula": f"Sparkline ({sl_info['type']})",
                        "value": f"Range: {sl_info['range']}",
                        "sparkline_type": sl_info["type"],
                        "sparkline_range": sl_info["range"],
                        "active_markers": active_markers if active_markers else "none",
                    })
                else:
                    cell_data.append({
                        "sheet": sn, "cell": c,
                        "type": "formula" if str(f).startswith("=") else "value",
                        "formula": str(f), "value": str(v),
                    })
            except Exception:
                pass

    if not cell_data:
        return {}, "", []

    prompt = f"""You are an Excel course instructor. Analyze the following answer key cells and generate a concise grading rubric.
Each cell entry includes a "type" field: "formula", "value", or "sparkline".
For sparkline cells, use the sparkline_type and sparkline_range fields to understand what is being tested.

Answer Key Data:
{json.dumps(cell_data, ensure_ascii=False, indent=2)}

Instructions:
- Group cells that test the SAME skill and criterion into one rubric item.
- Keep criterion names short (max 5 words).
- Keep common_mistakes brief (max 10 words).
- Suggested points should total ~100 across all items.

Return ONLY a valid JSON object in this exact format, no extra text:
{{
  "rubric_groups": [
    {{
      "cells": [{{"sheet": "Sheet1", "cell": "B2"}}, {{"sheet": "Sheet1", "cell": "B3"}}],
      "skill": "SUM Function",
      "criterion": "Correct SUM range",
      "points": 20,
      "common_mistakes": "Wrong range or hardcoded value"
    }}
  ],
  "task_summary": "One sentence describing what this lab tests"
}}"""

    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": "You are a professional Excel instructor. Return only valid JSON."},
                {"role": "user",   "content": prompt}
            ]
        )
        raw = response.choices[0].message.content.strip()
        raw = re.sub(r'^```json\s*', '', raw)
        raw = re.sub(r'\s*```$',     '', raw)
        rubric_data = json.loads(raw)

        rubric_map = {}
        for group in rubric_data.get("rubric_groups", []):
            for cell_ref in group.get("cells", []):
                key = (cell_ref["sheet"], cell_ref["cell"])
                rubric_map[key] = group

        return rubric_map, rubric_data.get("task_summary", ""), rubric_data.get("rubric_groups", [])
    except Exception as e:
        st.warning(f"Rubric generation failed: {e}")
        return {}, "", []

# ──────────────────────────────────────────────
# 11. PDF Report
# ──────────────────────────────────────────────
UNICODE_REPLACEMENTS = {
    '\u2022': '-', '\u2013': '-', '\u2014': '-',
    '\u2018': "'", '\u2019': "'", '\u201c': '"', '\u201d': '"',
    '\u2026': '...', '\u00b7': '-', '\u2192': '->',
    '\u2260': '!=', '\u2264': '<=', '\u2265': '>=',
}

def safe_text(text):
    if not text:
        return ""
    for char, replacement in UNICODE_REPLACEMENTS.items():
        text = text.replace(char, replacement)
    return text.encode('latin-1', 'replace').decode('latin-1')


def create_pdf_report(uid, score, errors, rubric_map=None):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 15, txt="IS 2010 Lab Diagnostic Report", ln=True, align='C')
    pdf.set_font("Arial", '', 12)
    pdf.cell(200, 10, txt=f"Student ID: {uid} | Final Score: {score}", ln=True, align='C')
    pdf.line(10, 35, 200, 35)
    pdf.ln(10)

    if not errors:
        pdf.set_font("Arial", 'B', 14)
        pdf.cell(200, 10, txt="Congratulations! Perfect Score.", ln=True, align='C')
    else:
        for i, err in enumerate(errors):
            pdf.set_font("Arial", 'B', 11)
            pdf.set_text_color(0, 0, 0)
            mode_label = err.get('Mode', 'Value Only')
            pdf.cell(0, 10, txt=safe_text(f"{i+1}. {err['Sheet']} / Cell {err['Cell']}  [{mode_label}]"), ln=True)

            pdf.set_font("Arial", '', 10)
            pdf.set_text_color(200, 0, 0)
            pdf.multi_cell(0, 7, txt=safe_text(f"Your Answer: {err['Stud_Ans']}"))
            pdf.set_text_color(0, 100, 0)
            pdf.multi_cell(0, 7, txt=safe_text(f"Expected Answer: {err['Prof_Ans']}"))

            if err.get('Issues'):
                pdf.set_text_color(150, 75, 0)
                for issue in err['Issues']:
                    pdf.multi_cell(0, 6, txt=safe_text(f"  - {issue}"))

            rubric_item = (rubric_map or {}).get((err['Sheet'], err['Cell']))
            if rubric_item:
                pdf.set_text_color(0, 0, 150)
                pdf.set_font("Arial", 'B', 10)
                rubric_str = f"[Rubric] {rubric_item['criterion']} ({rubric_item['points']}pts) -- {rubric_item.get('common_mistakes','')}"
                pdf.multi_cell(0, 7, txt=safe_text(rubric_str))

            rubric_ctx = rubric_item
            feedback = get_ai_feedback(err['Prof_Ans'], err['Stud_Ans'], issues=err.get('Issues'), rubric=rubric_ctx)
            pdf.set_fill_color(245, 245, 245)
            pdf.set_text_color(0, 50, 100)
            pdf.set_font("Arial", 'I', 10)
            pdf.multi_cell(0, 8, txt=safe_text(f"AI Feedback: {feedback or 'No feedback.'}"), border=1, fill=True)
            pdf.ln(5)

    return pdf.output(dest='S').encode('latin-1', 'ignore')

# ──────────────────────────────────────────────
# 12. Session State Init
# ──────────────────────────────────────────────
if 'grading_done' not in st.session_state:
    st.session_state.update({
        'grading_done':    False,
        'summary_df':      pd.DataFrame(),
        'all_wrongs_list': [],
        'total_questions': 0,
        'rubric_map':      {},
        'rubric_summary':  '',
        'rubric_groups':   [],
        'rubric_edited':   [],  # Professor's finalized rubric after editing
        'color_mode_map':  {},
    })

# ──────────────────────────────────────────────
# 13. UI: Color Guide
# ──────────────────────────────────────────────
with st.expander("Professor's Color Guide", expanded=True):
    cols = st.columns(10)
    st.warning(
        "※ **Crucial:** Use these 10 standard colors as **CELL SHADING (Fill Color)** to indicate graded cells. "
        "For each color, select the scoring mode below."
    )
    for idx, (name, hex_val) in enumerate(STANDARD_COLORS.items()):
        cols[idx].markdown(f"<div style='background-color:{hex_val}; height:20px; border-radius:3px;'></div>", unsafe_allow_html=True)
        cols[idx].caption(name)

# ──────────────────────────────────────────────
# 14. UI: Guidelines
# ──────────────────────────────────────────────
with st.expander("📘 Guided Lab Guidelines & Cautions", expanded=False):
    st.markdown("""
    #### 1. File Naming Convention (UNID Required)
    * **Student Files:** Filename must include your UID (e.g., **U1234567**).

    #### 2. Maintain Sheet Structure
    * Do not rename sheets or insert/delete rows or columns.

    #### 3. Precision in Formulas and Values
    * The system performs strict string comparison for formulas and numeric comparison for values.

    #### 4. Decimal Precision & Rounding Rules
    * Match the instructor's rounding method exactly (e.g., `=ROUND(calc, 2)`).

    #### 5. File Integrity
    * Only **.xlsx** supported. No password protection, no macros.
    """)

st.divider()

# ──────────────────────────────────────────────
# 15. File Upload
# ──────────────────────────────────────────────
c1, c2 = st.columns(2)
prof_file     = c1.file_uploader("1. Upload Professor's File", type=['xlsx'])
student_files = c2.file_uploader("2. Upload Student File(s)", type=['xlsx'], accept_multiple_files=True)

# ──────────────────────────────────────────────
# 16. Color → Scoring Mode Mapping UI
# ──────────────────────────────────────────────
if prof_file and student_files:
    st.subheader("🎨 Step 1: Assign Colors to Each Scoring Mode")
    st.caption("Select one color per scoring mode. Each color can only be used once.")

    NONE = "--- None ---"
    color_names = [NONE] + list(STANDARD_COLORS.keys())

    c1, c2, c3 = st.columns(3)
    with c3:
        st.markdown("**🔢🎨 Both**")
        both_color = st.selectbox("bc", color_names, key="sel_both", label_visibility="collapsed")

    restricted = {both_color} if both_color != NONE else set()
    remaining  = [c for c in color_names if c not in restricted]

    with c1:
        st.markdown("**🔢 Value Only**")
        value_color = st.selectbox("vc", remaining, key="sel_value", label_visibility="collapsed")
    with c2:
        fmt_opts_list = [c for c in remaining if c == NONE or c != value_color]
        st.markdown("**🎨 Format Only**")
        format_color = st.selectbox("fc", fmt_opts_list, key="sel_format", label_visibility="collapsed")

    # Color swatches
    sw1, sw2, sw3 = st.columns(3)
    for col, chosen in [(sw1, value_color), (sw2, format_color), (sw3, both_color)]:
        if chosen != NONE:
            col.markdown(
                f"<div style='background:{STANDARD_COLORS[chosen]};height:14px;border-radius:4px;'></div>",
                unsafe_allow_html=True
            )

    # Build color_mode_map
    color_mode_map = {}
    active_colors  = []
    for chosen, mode in [(value_color, "value_only"), (format_color, "format_only"), (both_color, "both")]:
        if chosen != NONE:
            hex_upper = STANDARD_COLORS[chosen].lstrip('#').upper()
            color_mode_map[hex_upper] = [mode]
            active_colors.append(chosen)

    if not active_colors:
        st.warning("⚠️ Please select at least one color.")

    # ──────────────────────────────────────────
    # 17. AI Rubric Generation + Customization
    # ──────────────────────────────────────────
    st.divider()
    st.subheader("🤖 Step 2: AI Rubric Auto-Generation (Optional)")

    if st.button("✨ Auto-Generate Rubric from Professor's File", disabled=(not active_colors)):
        with st.spinner("Analyzing formulas and generating rubric..."):
            prof_file.seek(0)
            p_bytes_r = prof_file.read()
            prof_file.seek(0)
            p_wb_f_r = load_workbook(io.BytesIO(p_bytes_r), data_only=False, read_only=False)
            p_wb_v_r = load_workbook(io.BytesIO(p_bytes_r), data_only=True,  read_only=False)
            p_sheet_map_r = build_sheet_xml_map(p_bytes_r)
            p_cache_r     = build_xml_cache(p_bytes_r, p_sheet_map_r)

            temp_check_map = {}
            for sn in p_wb_f_r.sheetnames:
                cells = [
                    cell.coordinate
                    for row in p_wb_f_r[sn].iter_rows()
                    for cell in row
                    if cell.fill and cell.fill.fill_type == 'solid'
                    and str(cell.fill.start_color.rgb)[-6:].upper() in color_mode_map
                ]
                if cells:
                    temp_check_map[sn] = cells

            rubric_map, rubric_summary, rubric_groups = generate_rubric(
                temp_check_map, p_wb_f_r, p_wb_v_r, p_cache=p_cache_r
            )
            st.session_state['rubric_map']     = rubric_map
            st.session_state['rubric_summary'] = rubric_summary
            st.session_state['rubric_groups']  = rubric_groups
            st.session_state['rubric_edited']  = [dict(g) for g in rubric_groups]

    # ── 루브릭 편집 UI ────────────────────────
    if st.session_state.get('rubric_summary'):
        st.success(f"📝 {st.session_state['rubric_summary']}")
        edited = st.session_state.get('rubric_edited', [])

        if edited:
            st.markdown("#### ✏️ Customize Rubric")
            st.caption("Set Total Score to auto-distribute points by ratio, or adjust each item individually.")

            # ── Total Score 입력 ──
            ai_total = sum(g.get('points', 0) for g in edited)
            tc1, tc2 = st.columns([1, 4])
            with tc1:
                new_total = st.number_input(
                    "🎯 Total Score",
                    min_value=0.0, max_value=1000.0,
                    value=float(ai_total),
                    step=0.5,
                    key="rubric_total",
                    help="Change this to redistribute all item points by ratio automatically."
                )
            with tc2:
                if ai_total > 0 and abs(new_total - ai_total) > 0.01:
                    st.info(f"💡 Points will be redistributed: {ai_total:.1f}pts → {new_total:.1f}pts (ratio kept)")

            # Total 변경 시 비율로 재계산 (UI 렌더링 전에 적용)
            display_groups = []
            for g in edited:
                g_copy = dict(g)
                if ai_total > 0 and abs(new_total - ai_total) > 0.01:
                    g_copy['points'] = round(g['points'] * (new_total / ai_total), 1)
                display_groups.append(g_copy)

            # ── 항목별 편집 ──
            st.markdown("---")
            updated_groups = []
            for idx, g in enumerate(display_groups):
                cells_str = ", ".join(
                    f"{c['sheet']}!{c['cell']}" if len(set(cr['sheet'] for cr in g['cells'])) > 1
                    else c['cell']
                    for c in g['cells']
                )
                row_l, row_r = st.columns([5, 1])
                with row_l:
                    st.markdown(f"**{idx+1}. {g['skill']}** &nbsp;·&nbsp; {g['criterion']}")
                    st.caption(f"📍 {cells_str}　　⚠️ {g.get('common_mistakes', '')}")
                with row_r:
                    pts = st.number_input(
                        f"pts_{idx}",
                        min_value=0.0, max_value=500.0,
                        value=float(g['points']),
                        step=0.5,
                        key=f"rubric_pts_{idx}",
                        label_visibility="collapsed"
                    )
                g_updated = dict(g)
                g_updated['points'] = pts
                updated_groups.append(g_updated)

            # ── 합계 & Apply 버튼 ──
            current_total = sum(g['points'] for g in updated_groups)
            m1, m2 = st.columns([1, 4])
            with m1:
                delta = current_total - new_total
                st.metric(
                    "Current Total",
                    f"{current_total:.1f} pts",
                    delta=f"{delta:+.1f}" if abs(delta) > 0.01 else None,
                    delta_color="inverse" if delta > 0.01 else ("off" if delta < -0.01 else "normal")
                )
            with m2:
                if st.button("✅ Apply & Lock Rubric", width="stretch"):
                    new_rubric_map = {}
                    for g in updated_groups:
                        for cell_ref in g.get('cells', []):
                            key = (cell_ref['sheet'], cell_ref['cell'])
                            new_rubric_map[key] = g
                    st.session_state['rubric_map']    = new_rubric_map
                    st.session_state['rubric_edited'] = updated_groups
                    st.success(f"✅ Rubric locked! Total: {current_total:.1f} pts · {len(updated_groups)} criteria")

    # ──────────────────────────────────────────
    # 18. Grading Button
    # ──────────────────────────────────────────
    st.divider()
    st.subheader("📊 Step 3: Start Grading")

    # --- Format Options ---
    with st.expander("🎨 Format Grading Options (applies to Format Only / Both modes)", expanded=False):
        st.caption("Select which formatting attributes to check. Unchecked items are ignored during grading.")
        fo1, fo2, fo3 = st.columns(3)
        fmt_options = {
            "number_format": fo1.checkbox("Number Format",  value=True),
            "font_name":     fo1.checkbox("Font Name",      value=False),
            "font_bold":     fo1.checkbox("Bold",           value=True),
            "font_color":    fo2.checkbox("Font Color",     value=True),
            "font_size":     fo2.checkbox("Font Size",      value=True),
            "border":        fo2.checkbox("Border",         value=True),
            "alignment":     fo3.checkbox("Alignment",      value=False),
            "wrap_text":     fo3.checkbox("Wrap Text",      value=False),
            "cell_size":     fo3.checkbox("Cell Size (Row Height / Column Width)", value=False),
            "merge":         fo3.checkbox("Cell Merge",     value=False),
        }

    # --- Pre-check Settings ---
    with st.expander("⚙️ Pre-check Settings (File Name Validation)", expanded=True):
        pc_col1, pc_col2 = st.columns(2)
        with pc_col1:
            st.markdown("**📄 File Name Format**")
            filename_format = st.text_input(
                "Expected filename format (use [UID] as placeholder for student ID)",
                value="[UID]_Lab_1.xlsx",
                help="e.g. '[UID]_Lab_1.xlsx' → matches 'u1234567_Lab_1.xlsx'\nLeave blank to skip."
            )
        with pc_col2:
            st.markdown("**📉 Penalty Points**")
            penalty_filename = st.number_input(
                "File name mismatch penalty (pts)",
                min_value=0.0, max_value=50.0, value=0.0, step=0.5,
                help="Set to 0 to show warning only without deducting points"
            )
        st.caption("💡 Set penalty to 0 to show warnings only without affecting the score.")

    if st.button("🚀 Start Grading Process", width='stretch', disabled=(not active_colors)):

        st.session_state.update({
            'grading_done':    False,
            'summary_df':      pd.DataFrame(),
            'all_wrongs_list': [],
            'total_questions': 0,
            'zip_data':        None,
            # Preserve rubric state so edits survive a re-grade
            'rubric_map':      st.session_state.get('rubric_map', {}),
            'rubric_summary':  st.session_state.get('rubric_summary', ''),
            'rubric_groups':   st.session_state.get('rubric_groups', []),
            'rubric_edited':   st.session_state.get('rubric_edited', []),
        })

        prof_file.seek(0)
        p_bytes = prof_file.read()
        prof_file.seek(0)

        p_sheet_map = build_sheet_xml_map(p_bytes)
        p_cache     = build_xml_cache(p_bytes, p_sheet_map)
        p_wb_f = load_workbook(io.BytesIO(p_bytes), data_only=False, read_only=False)
        p_wb_v = load_workbook(io.BytesIO(p_bytes), data_only=True,  read_only=False)

        check_map = {}
        for sn in p_wb_f.sheetnames:
            entries = []
            for row in p_wb_f[sn].iter_rows():
                for cell in row:
                    if cell.fill and cell.fill.fill_type == 'solid':
                        rgb = str(cell.fill.start_color.rgb)[-6:].upper()
                        if rgb in color_mode_map:
                            entries.append((cell.coordinate, color_mode_map[rgb]))
            if entries:
                check_map[sn] = entries

        total_qs = sum(len(v) for v in check_map.values())
        summary_results = []
        all_wrongs = []
        uid_re = re.compile(r'[uU]\d{7}')
        progress_bar = st.progress(0)

        def build_filename_regex(fmt):
            if not fmt:
                return None
            escaped = re.escape(fmt)
            escaped = escaped.replace(r'\[UID\]', r'[uU]\d{7}')
            return re.compile(escaped, re.IGNORECASE)

        filename_regex  = build_filename_regex(filename_format)

        for i, s_file in enumerate(student_files):
            with st.status(f"Grading {s_file.name}...", expanded=True) as status:
                s_bytes = s_file.read()
                s_sheet_map = build_sheet_xml_map(s_bytes)
                s_cache     = build_xml_cache(s_bytes, s_sheet_map)
                s_wb_f = load_workbook(io.BytesIO(s_bytes), data_only=False, read_only=False)
                s_wb_v = load_workbook(io.BytesIO(s_bytes), data_only=True,  read_only=False)
                uid = uid_re.search(s_file.name).group() if uid_re.search(s_file.name) else s_file.name

                precheck_warnings = []
                penalty = 0.0

                if filename_regex:
                    if not filename_regex.search(s_file.name):
                        msg = f"File name mismatch: expected '{filename_format}', got '{s_file.name}'"
                        precheck_warnings.append(msg)
                        penalty += penalty_filename
                        st.warning(f"⚠️ {msg}" + (f" (-{penalty_filename}pts)" if penalty_filename > 0 else ""))
                    else:
                        st.success(f"✅ File name OK: {s_file.name}")

                correct = 0
                for sn, cell_entries in check_map.items():
                    if sn not in s_wb_f.sheetnames:
                        continue
                    for (c, modes) in cell_entries:
                        cell_correct = True
                        cell_issues  = []
                        stud_ans_str = "N/A"
                        prof_ans_str = "N/A"

                        for mode in modes:
                            is_correct, stud_ans, prof_ans, issues = grade_cell(
                                mode, p_wb_f, p_wb_v, s_wb_f, s_wb_v, p_cache, s_cache, sn, c,
                                fmt_options=fmt_options
                            )
                            stud_ans_str = stud_ans
                            prof_ans_str = prof_ans
                            if not is_correct:
                                cell_correct = False
                                for iss in issues:
                                    cell_issues.append(f"[{SCORING_LABELS[mode]}] {iss}")

                        if cell_correct:
                            correct += 1
                        else:
                            all_wrongs.append({
                                "UnID":     uid,
                                "Sheet":    sn,
                                "Cell":     c,
                                "Mode":     " + ".join([SCORING_LABELS[m] for m in modes]),
                                "Stud_Ans": stud_ans_str,
                                "Prof_Ans": prof_ans_str,
                                "Issues":   cell_issues,
                            })

                status.update(label=f"✅ {uid} Done!", state="complete", expanded=False)

            try:
                s_wb_f.close()
                s_wb_v.close()
            except Exception:
                pass

            # Calculate rubric-scaled score
            rubric_edited = st.session_state.get('rubric_edited', [])
            rubric_total  = sum(g.get('points', 0) for g in rubric_edited) if rubric_edited else 0
            if rubric_total > 0 and total_qs > 0:
                rubric_score     = round((correct / total_qs) * rubric_total, 2)
                rubric_score_str = f"{rubric_score} / {rubric_total}"
                total_score      = round(max(rubric_score - penalty, 0), 2)
                total_score_str  = f"{total_score} / {rubric_total}"
            else:
                rubric_score_str = "N/A (no rubric)"
                total_score_str  = "N/A (no rubric)"

            summary_results.append({
                "UnID":          uid,
                "Raw Score":     f"{correct}/{total_qs}",
                "Rubric Score":  rubric_score_str,
                "Penalty":       f"-{penalty}pts" if penalty > 0 else "—",
                "Total Score":   total_score_str,
                "Raw":           correct,
                "Warnings":      " | ".join(precheck_warnings) if precheck_warnings else "✅ OK",
            })
            progress_bar.progress((i + 1) / len(student_files))

        try:
            p_wb_f.close()
            p_wb_v.close()
        except Exception:
            pass

        st.session_state.update({
            'summary_df': pd.DataFrame(summary_results).sort_values("Raw", ascending=False),
            'all_wrongs_list': all_wrongs,
            'grading_done': True,
            'total_questions': total_qs,
            'color_mode_map': color_mode_map,
        })
        st.rerun()

# ──────────────────────────────────────────────
# 19. Results Display
# ──────────────────────────────────────────────
if st.session_state['grading_done']:
    df_all_errors = pd.DataFrame(st.session_state['all_wrongs_list'])

    st.divider()
    col_chart, col_table = st.columns([1, 1])
    with col_chart:
        st.subheader("📊 Class Error Analysis")
        if not df_all_errors.empty:
            st.bar_chart(df_all_errors['Cell'].value_counts())
        else:
            st.success("No errors detected! 🎉")
    with col_table:
        st.subheader("📋 Score Summary")
        st.dataframe(
            st.session_state['summary_df'].drop(columns=["Raw"], errors='ignore'),
            use_container_width=True,
            hide_index=True
        )

    if not df_all_errors.empty:
        st.divider()
        st.subheader("🔍 Error Detail by Scoring Mode")
        if 'Mode' in df_all_errors.columns:
            for mode_label in df_all_errors['Mode'].unique():
                mode_df = df_all_errors[df_all_errors['Mode'] == mode_label]
                if not mode_df.empty:
                    with st.expander(f"{mode_label} Errors ({len(mode_df)})"):
                        display_df = mode_df[['UnID', 'Sheet', 'Cell', 'Stud_Ans', 'Prof_Ans']].copy()
                        st.dataframe(display_df, use_container_width=True, hide_index=True)
        else:
            st.info("No mode information available.")

    st.divider()
    st.subheader("📥 Download Reports")
    col_dl1, col_dl2 = st.columns(2)

    with col_dl1:
        xlsx_report = io.BytesIO()
        with pd.ExcelWriter(xlsx_report, engine='openpyxl') as writer:
            st.session_state['summary_df'].drop(columns=['Raw'], errors='ignore').to_excel(writer, index=False, sheet_name="Summary")
            if not df_all_errors.empty:
                export_df = df_all_errors.copy()
                if 'Issues' in export_df.columns:
                    export_df['Issues'] = export_df['Issues'].apply(lambda x: "; ".join(x) if isinstance(x, list) else x)
                export_df.to_excel(writer, index=False, sheet_name="All_Errors")
            # Export finalized rubric if available
            if st.session_state.get('rubric_edited'):
                rubric_export = pd.DataFrame([
                    {
                        "Cells":          ", ".join(f"{c['sheet']}!{c['cell']}" for c in g['cells']),
                        "Skill":          g["skill"],
                        "Criterion":      g["criterion"],
                        "Points":         g["points"],
                        "Common Mistakes": g.get("common_mistakes", ""),
                    }
                    for g in st.session_state['rubric_edited']
                ])
                rubric_export.to_excel(writer, index=False, sheet_name="Rubric")
        st.download_button(
            "📊 Download Excel Summary",
            xlsx_report.getvalue(),
            "IS2010_Results.xlsx",
            width='stretch'
        )

    with col_dl2:
        if st.button("📄 Step 1: Generate AI Reports", width="stretch"):
            zip_buffer = io.BytesIO()
            rubric_map = st.session_state.get('rubric_map', {})
            failed = []
            with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zf:
                for uid in st.session_state['summary_df']['UnID'].tolist():
                    with st.spinner(f"Generating AI Feedback for {uid}..."):
                        try:
                            s_errs  = [e for e in st.session_state['all_wrongs_list'] if e['UnID'] == uid]
                            s_score = st.session_state['summary_df'][
                                st.session_state['summary_df']['UnID'] == uid
                            ]['Raw Score'].values[0]
                            pdf_data = create_pdf_report(uid, s_score, s_errs, rubric_map=rubric_map)
                            zf.writestr(f"Report_{uid}.pdf", pdf_data)
                        except Exception as e:
                            failed.append(uid)
                            st.warning(f"⚠️ Failed to generate report for {uid}: {e}")
            st.session_state['zip_data'] = zip_buffer.getvalue()
            if failed:
                st.warning(f"⚠️ {len(failed)} report(s) failed: {', '.join(failed)}")
            else:
                st.success("✅ All reports generated!")

        if 'zip_data' in st.session_state and st.session_state['zip_data'] is not None:
            st.download_button(
                label="📥 Download All PDF Reports (ZIP)",
                data=st.session_state['zip_data'],
                file_name="Student_Reports.zip",
                mime="application/zip",
                width='stretch'
            )
else:
    st.info("Upload files and configure scoring modes to begin.")
